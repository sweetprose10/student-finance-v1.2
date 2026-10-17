from datetime import datetime
from functools import wraps
from io import BytesIO
import os
from pathlib import Path

from flask import Flask, flash, redirect, render_template, request, send_file, session, url_for
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import database

try:
    from dotenv import load_dotenv
except ImportError:
    load_dotenv = None

if load_dotenv:
    load_dotenv(Path(__file__).with_name(".env"))

try:
    from authlib.integrations.flask_client import OAuth
except ImportError:
    OAuth = None

app = Flask(__name__)
app.secret_key = os.environ.get("STUDENTFINANCE_SECRET_KEY", "dev-only-change-this-secret")
app.config.update(
    SESSION_COOKIE_HTTPONLY=True,
    SESSION_COOKIE_SAMESITE="Lax",
    SESSION_COOKIE_SECURE=os.environ.get("FLASK_SECURE_COOKIES", "0") == "1",
)

google_enabled = bool(OAuth and os.environ.get("GOOGLE_CLIENT_ID") and os.environ.get("GOOGLE_CLIENT_SECRET"))
oauth = OAuth(app) if google_enabled else None
if google_enabled:
    oauth.register(
        name="google",
        client_id=os.environ["GOOGLE_CLIENT_ID"],
        client_secret=os.environ["GOOGLE_CLIENT_SECRET"],
        server_metadata_url="https://accounts.google.com/.well-known/openid-configuration",
        client_kwargs={"scope": "openid email profile"},
    )


@app.context_processor
def inject_auth_config():
    return {"google_enabled": google_enabled}


def login_required(view):
    @wraps(view)
    def wrapped_view(*args, **kwargs):
        if not session.get("user_id"):
            return redirect(url_for("login"))
        return view(*args, **kwargs)

    return wrapped_view


def admin_required(view):
    @wraps(view)
    def wrapped_view(*args, **kwargs):
        if session.get("role") != "admin":
            flash("Halaman ini khusus admin developer.", "error")
            return redirect(url_for("dashboard"))
        return view(*args, **kwargs)

    return wrapped_view


def current_user_id():
    return int(session["user_id"])


def parse_number(field):
    try:
        return int(request.form.get(field, "0"))
    except ValueError:
        return 0


@app.route("/")
@login_required
def dashboard():
    user_id = current_user_id()
    return render_template(
        "index.html",
        saldo=database.saldo(user_id),
        pemasukan=database.total_pemasukan(user_id),
        pengeluaran=database.total_pengeluaran(user_id),
        transaksi=database.semua_transaksi(user_id),
    )


@app.route("/pemasukan")
@login_required
def pemasukan():
    user_id = current_user_id()
    cari = request.args.get("cari")
    data = database.cari_transaksi(user_id, cari) if cari else database.semua_transaksi(user_id)
    return render_template("pemasukan.html", data=data)


@app.route("/tambah_pemasukan", methods=["POST"])
@login_required
def tambah_pemasukan():
    nominal = parse_number("nominal")
    if nominal <= 0:
        flash("Nominal pemasukan harus lebih dari 0.", "error")
        return redirect(url_for("pemasukan"))
    database.tambah_transaksi(current_user_id(), "Pemasukan", request.form.get("kategori", "").strip(), nominal, request.form.get("keterangan", "").strip(), datetime.now().strftime("%d/%m/%Y"))
    flash("Pemasukan berhasil ditambahkan.")
    return redirect(url_for("dashboard"))


@app.route("/pengeluaran")
@login_required
def pengeluaran():
    return render_template("pengeluaran.html", data=database.semua_transaksi(current_user_id()))


@app.route("/tambah_pengeluaran", methods=["POST"])
@login_required
def tambah_pengeluaran():
    nominal = parse_number("nominal")
    if nominal <= 0:
        flash("Nominal pengeluaran harus lebih dari 0.", "error")
        return redirect(url_for("pengeluaran"))
    database.tambah_transaksi(current_user_id(), "Pengeluaran", request.form.get("kategori", "").strip(), nominal, request.form.get("keterangan", "").strip(), datetime.now().strftime("%d/%m/%Y"))
    flash("Pengeluaran berhasil ditambahkan.")
    return redirect(url_for("dashboard"))


@app.route("/hapus/<int:id>")
@login_required
def hapus(id):
    database.hapus_transaksi(current_user_id(), id)
    flash("Transaksi berhasil dihapus.")
    return redirect(url_for("dashboard"))


@app.route("/edit/<int:id>")
@login_required
def edit(id):
    return render_template("edit.html", data=database.transaksi_by_id(current_user_id(), id))


@app.route("/update/<int:id>", methods=["POST"])
@login_required
def update(id):
    nominal = parse_number("nominal")
    if nominal <= 0:
        flash("Nominal harus lebih dari 0.", "error")
        return redirect(url_for("edit", id=id))
    database.update_transaksi(current_user_id(), id, request.form.get("kategori", "").strip(), nominal, request.form.get("keterangan", "").strip())
    flash("Transaksi berhasil diperbarui.")
    return redirect(url_for("dashboard"))


@app.route("/wishlist")
@login_required
def wishlist():
    return render_template("wishlist.html", data=database.semua_wishlist(current_user_id()))


@app.route("/tambah_wishlist", methods=["POST"])
@login_required
def tambah_wishlist():
    harga = parse_number("harga")
    if harga <= 0:
        flash("Harga wishlist harus lebih dari 0.", "error")
        return redirect(url_for("wishlist"))
    database.tambah_wishlist(current_user_id(), request.form.get("nama", "").strip(), harga, request.form.get("prioritas", "Normal"))
    flash("Wishlist berhasil ditambahkan.")
    return redirect(url_for("wishlist"))


@app.route("/hapus_wishlist/<int:id>")
@login_required
def hapus_wishlist(id):
    database.hapus_wishlist(current_user_id(), id)
    flash("Wishlist berhasil dihapus.")
    return redirect(url_for("wishlist"))


@app.route("/target")
@login_required
def target():
    return render_template("target.html", targets=database.semua_target(current_user_id()))


@app.route("/tambah_target", methods=["POST"])
@login_required
def tambah_target():
    nominal = parse_number("nominal")
    if nominal <= 0 or not request.form.get("nama", "").strip():
        flash("Nama dan nominal target wajib diisi.", "error")
        return redirect(url_for("target"))
    database.tambah_target(current_user_id(), request.form["nama"].strip(), nominal, request.form.get("tenggat", "").strip())
    flash("Target tabungan berhasil dibuat.")
    return redirect(url_for("target"))


@app.route("/target/<int:id>/progress", methods=["POST"])
@login_required
def update_target_progress(id):
    terkumpul = parse_number("terkumpul")
    if terkumpul < 0:
        flash("Jumlah terkumpul tidak boleh negatif.", "error")
    else:
        database.update_target_progress(current_user_id(), id, terkumpul)
        flash("Progress target diperbarui.")
    return redirect(url_for("target"))


@app.route("/hapus_target/<int:id>")
@login_required
def hapus_target(id):
    database.hapus_target(current_user_id(), id)
    flash("Target berhasil dihapus.")
    return redirect(url_for("target"))


@app.route("/statistik")
@login_required
def statistik():
    user_id = current_user_id()
    return render_template("statistik.html", pemasukan=database.total_pemasukan(user_id), pengeluaran=database.total_pengeluaran(user_id))


@app.route("/profil")
@login_required
def profil():
    return render_template("profil.html", user=database.user_by_id(current_user_id()))


@app.route("/hapus-akun", methods=["POST"])
@login_required
def hapus_akun():
    user_id = current_user_id()
    if session.get("role") == "admin":
        flash("Akun admin developer tidak dapat dihapus.", "error")
        return redirect(url_for("profil"))
    database.delete_user(user_id)
    session.clear()
    return redirect(url_for("login"))


@app.route("/admin/developer")
@login_required
@admin_required
def developer_admin():
    return render_template("developer.html", users=database.all_users())


@app.route("/export_excel")
@login_required
def export_excel():
    data = database.semua_transaksi(current_user_id())
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Transaksi"
    worksheet.sheet_view.showGridLines = False
    worksheet.merge_cells("A1:E1")
    worksheet["A1"] = "LAPORAN KEUANGAN STUDENTFINANCE"
    worksheet["A1"].font = Font(size=16, bold=True, color="FFFFFF")
    worksheet["A1"].fill = PatternFill("solid", fgColor="172B4D")
    worksheet["A1"].alignment = Alignment(horizontal="center")
    worksheet.merge_cells("A2:E2")
    worksheet["A2"] = f"Dibuat pada {datetime.now().strftime('%d %B %Y, %H:%M')}"
    worksheet["A2"].font = Font(italic=True, color="64748B")
    headers = ["Jenis", "Kategori", "Nominal", "Keterangan", "Tanggal"]
    for column, header in enumerate(headers, 1):
        cell = worksheet.cell(4, column, header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2F80ED")
        cell.alignment = Alignment(horizontal="center")
    for row, item in enumerate(data, 5):
        values = [item[2], item[3], item[4], item[5] or "-", item[6]]
        for column, value in enumerate(values, 1):
            worksheet.cell(row, column, value)
        worksheet.cell(row, 3).number_format = '"Rp" #,##0'
        worksheet.cell(row, 1).font = Font(color="188B5A" if item[2] == "Pemasukan" else "D95D39", bold=True)
    for column, width in enumerate([18, 20, 18, 36, 18], 1):
        worksheet.column_dimensions[get_column_letter(column)].width = width
    thin_border = Border(bottom=Side(style="thin", color="E2E8F0"))
    for row in worksheet.iter_rows(min_row=4, max_row=max(4, worksheet.max_row), min_col=1, max_col=5):
        for cell in row:
            cell.border = thin_border
    worksheet.auto_filter.ref = f"A4:E{max(4, worksheet.max_row)}"
    worksheet.freeze_panes = "A5"
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return send_file(output, as_attachment=True, download_name="laporan-keuangan.xlsx", mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


@app.route("/login", methods=["GET", "POST"])
def login():
    if request.method == "POST":
        user = database.get_user(request.form.get("username", "").strip(), request.form.get("password", ""))
        if user:
            session["user_id"] = user["id"]
            session["username"] = user["username"]
            session["role"] = user["role"]
            return redirect(url_for("dashboard"))
        return render_template("login.html", error="Username atau password salah.")
    return render_template("login.html")


@app.route("/login/google")
def google_login():
    if not google_enabled:
        flash("Login Google belum dikonfigurasi oleh administrator.", "error")
        return redirect(url_for("login"))
    redirect_uri = url_for("google_callback", _external=True)
    return oauth.google.authorize_redirect(redirect_uri)


@app.route("/login/google/callback")
def google_callback():
    if not google_enabled:
        return redirect(url_for("login"))
    try:
        token = oauth.google.authorize_access_token()
        profile = token.get("userinfo") or oauth.google.userinfo()
        google_sub = profile.get("sub")
        email = profile.get("email")
        if not google_sub or not email:
            flash("Google tidak mengirim identitas akun yang diperlukan.", "error")
            return redirect(url_for("login"))
        user = database.user_by_google_sub(google_sub)
        if not user:
            username = email.split("@", 1)[0][:24]
            user_id = database.create_google_user(username, google_sub)
            if not user_id:
                user_id = database.create_google_user(f"{username}_{google_sub[-6:]}", google_sub)
            user = database.user_by_id(user_id) if user_id else None
        if not user:
            flash("Akun Google tidak dapat dibuat.", "error")
            return redirect(url_for("login"))
        session["user_id"] = user["id"]
        session["username"] = user["username"]
        session["role"] = user["role"]
        return redirect(url_for("dashboard"))
    except Exception:
        app.logger.exception("Google OAuth callback failed")
        flash("Login Google gagal. Silakan coba lagi.", "error")
        return redirect(url_for("login"))


@app.route("/daftar", methods=["GET", "POST"])
def daftar():
    if request.method == "POST":
        username = request.form.get("username", "").strip()
        password = request.form.get("password", "")
        confirm = request.form.get("confirm", "")
        if len(username) < 3 or len(password) < 5:
            return render_template("register.html", error="Username minimal 3 karakter dan password minimal 5 karakter.")
        if password != confirm:
            return render_template("register.html", error="Konfirmasi password tidak cocok.")
        user_id = database.create_user(username, password)
        if not user_id:
            return render_template("register.html", error="Username sudah digunakan.")
        session["user_id"] = user_id
        session["username"] = username
        session["role"] = "user"
        return redirect(url_for("dashboard"))
    return render_template("register.html")


@app.route("/logout")
def logout():
    session.clear()
    return redirect(url_for("login"))


if __name__ == "__main__":
    app.run(debug=True)
